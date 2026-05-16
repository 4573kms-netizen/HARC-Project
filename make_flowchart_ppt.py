"""
HARC Simulator Flowchart — 가로형 PPT (13.33 x 7.5 inch, native shapes)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
from pptx.oxml.ns import qn
from lxml import etree

OUT = r'C:\Users\4573k\Desktop\HARC_simulation_Claude\HARC_Flowchart.pptx'

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def rgb(r, g, b): return RGBColor(r, g, b)
C = {
    's1':  rgb(236,236,236),
    's2':  rgb(212,212,212),
    's3':  rgb(180,180,180),
    's4':  rgb(148,148,148),
    's5':  rgb(115,115,115),
    's6':  rgb(225,225,225),
    'ann': rgb(250,250,250),
    'circ':rgb(46, 46, 46),
    'bdr': rgb(80, 80, 80),
    'dark':rgb(16, 16, 16),
    'lite':rgb(245,245,245),
    'loop':rgb(105,105,105),
    'abdr':rgb(175,175,175),
}

# ── 레이아웃 ──────────────────────────────────────────────
cx  = 4.8    # 메인 박스 중심 x (왼쪽 루프 라벨 공간 확보)
bw  = 6.4    # 박스 너비
bh  = 0.62   # 박스 높이
bh5 = 0.74   # STEP5 높이
rc  = 0.27   # START/END 원 반지름
axc = 11.2   # 어노테이션 중심 x
aw  = 4.0    # 어노테이션 너비

# Y 위치 — draw_flowchart.py 와 동일
ys = 0.55
y1 = 1.25
y2 = 1.99
y3 = 2.73
y4 = 3.47
y5 = 4.27
y6 = 5.07
ye = 5.77

# ══════════════════════════════════════════════════════════
# 헬퍼 함수
# ══════════════════════════════════════════════════════════
def rbox(cx_i, cy_i, w, h, title, sub, bg, tc, ft=14, fs=11):
    sp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(cx_i-w/2), Inches(cy_i-h/2), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = bg
    sp.line.color.rgb = C['bdr']; sp.line.width = Pt(1.5)
    tf = sp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER; p.space_before = Pt(4)
    r = p.add_run(); r.text = title
    r.font.bold = True; r.font.size = Pt(ft); r.font.color.rgb = tc
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER; p2.space_after = Pt(4)
    r2 = p2.add_run(); r2.text = sub
    r2.font.size = Pt(fs); r2.font.color.rgb = tc

def oval(cx_i, cy_i, r, bg, text, tc, fs=14):
    sp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(cx_i-r), Inches(cy_i-r), Inches(r*2), Inches(r*2))
    sp.fill.solid(); sp.fill.fore_color.rgb = bg
    sp.line.color.rgb = C['bdr']; sp.line.width = Pt(2.0)
    tf = sp.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r_ = tf.paragraphs[0].add_run(); r_.text = text
    r_.font.bold = True; r_.font.size = Pt(fs); r_.font.color.rgb = tc

def seg(x1, y1_, x2, y2_, col=None, lw=2.0):
    col = col or C['bdr']
    cn = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(x1), Inches(y1_), Inches(x2), Inches(y2_))
    cn.line.color.rgb = col; cn.line.width = Pt(lw)
    ln = cn.line._ln
    for t in [qn('a:headEnd'), qn('a:tailEnd')]:
        for e in ln.findall(t): ln.remove(e)

def arr(x1, y1_, x2, y2_, col=None, lw=2.0):
    """(x1,y1) → (x2,y2) 방향으로 화살촉 (tailEnd = 도착점)"""
    col = col or C['bdr']
    cn = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(x1), Inches(y1_), Inches(x2), Inches(y2_))
    cn.line.color.rgb = col; cn.line.width = Pt(lw)
    ln = cn.line._ln
    for t in [qn('a:headEnd'), qn('a:tailEnd')]:
        for e in ln.findall(t): ln.remove(e)
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'arrow'); tail.set('w', 'med'); tail.set('len', 'med')

def dseg(x1, y1_, x2, y2_, col=None, lw=0.9):
    """점선 연결 (화살촉 없음)"""
    col = col or C['abdr']
    cn = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(x1), Inches(y1_), Inches(x2), Inches(y2_))
    cn.line.color.rgb = col; cn.line.width = Pt(lw)
    ln = cn.line._ln
    for t in [qn('a:headEnd'), qn('a:tailEnd')]:
        for e in ln.findall(t): ln.remove(e)
    pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', 'dash')

def annbox(cx_i, cy_i, w, h, lines, fs=11):
    sp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(cx_i-w/2), Inches(cy_i-h/2), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = C['ann']
    sp.line.color.rgb = C['abdr']; sp.line.width = Pt(1.0)
    ln = sp.line._ln
    pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', 'dash')
    tf = sp.text_frame; tf.word_wrap = True
    for i, txt in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        if i == 0: p.space_before = Pt(3)
        if i == len(lines)-1: p.space_after = Pt(3)
        r_ = p.add_run(); r_.text = txt
        r_.font.size = Pt(fs); r_.font.color.rgb = C['dark']

def lbl(x, y, w, h, text, fs=11, bold=False, col=None, italic=False, align='center'):
    col = col or C['dark']
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if align == 'center' else PP_ALIGN.LEFT
    r_ = p.add_run(); r_.text = text
    r_.font.size = Pt(fs); r_.font.bold = bold
    r_.font.italic = italic; r_.font.color.rgb = col

# ══════════════════════════════════════════════════════════
# 제목
# ══════════════════════════════════════════════════════════
lbl(0.2, 0.04, 12.9, 0.36,
    'HARC Etch Simulator  —  Simulation Flow',
    fs=20, bold=True, col=C['dark'], align='center')

# ══════════════════════════════════════════════════════════
# 메인 박스
# ══════════════════════════════════════════════════════════
oval(cx, ys, rc, C['circ'], 'START', C['lite'])

rbox(cx, y1, bw, bh,
     'STEP 1  |  공정 조건 입력',
     'CF4/Ar 유량  ·  V_bias  ·  Source Power  ·  Pressure  ·  식각 시간',
     C['s1'], C['dark'])

rbox(cx, y2, bw, bh,
     'STEP 2  |  0-D 플라즈마 모델',
     '표면 플럭스 계산  (Γ_F, Γ_CFx, Γ_ion)  +  평균 이온 에너지',
     C['s2'], C['dark'])

rbox(cx, y3, bw, bh,
     'STEP 3  |  수송 모델',
     '플럭스가 홀 내부 깊이에 따라 얼마나 감쇠하는지 계산',
     C['s3'], C['dark'])

rbox(cx, y4, bw, bh,
     'STEP 4  |  표면 반응 모델',
     '수직 식각 속도  ·  측면 식각 속도  ·  스퍼터링 계산',
     C['s4'], C['lite'])

rbox(cx, y5, bw, bh5,
     'STEP 5  |  프로파일 시간 적분  (Δt 루프)',
     '깊이 증가  ·  CD 프로파일  ·  마스크 개구부  ·  폴리머 업데이트',
     C['s5'], C['lite'])

rbox(cx, y6, bw, bh,
     'STEP 6  |  출력',
     'Depth  ·  CD_top  ·  CD_bot  ·  Aspect Ratio  ·  Taper  ·  Bowing',
     C['s6'], C['dark'])

oval(cx, ye, rc, C['circ'], 'END', C['lite'])

# ══════════════════════════════════════════════════════════
# Δt 루프 (왼쪽)  STEP5 하단 → 왼쪽 → 위로 → STEP3 화살촉
# ══════════════════════════════════════════════════════════
lx    = cx - bw/2 - 0.55   # = 4.8 - 3.2 - 0.55 = 1.05
t_top = y3                  # STEP3 중심 y = 2.73
t_bot = y5 + bh5 / 2       # STEP5 하단 = 4.64

seg(cx - bw/2, t_bot, lx, t_bot, col=C['loop'])   # 하단 수평선
seg(lx, t_bot, lx, t_top,        col=C['loop'])    # 수직선 (아래→위)
arr(lx, t_top, cx - bw/2, t_top, col=C['loop'])   # STEP3 방향 화살촉

# 루프 라벨 — lx 왼쪽 공간 (x: 0.05 ~ lx-0.08 ≈ 0.97)
label_w = lx - 0.08   # ≈ 0.97 인치
mid_y   = (t_top + t_bot) / 2   # ≈ 3.685

lbl(0.05, mid_y - 0.32, label_w, 0.38,
    'STEP 5 → STEP 3\n으로 반복',
    fs=10, bold=True, col=C['loop'], align='center')

lbl(0.05, mid_y + 0.08, label_w, 0.28,
    '(Δt마다 수렴할 때까지)',
    fs=9, italic=True, col=C['loop'], align='center')

# ══════════════════════════════════════════════════════════
# 어노테이션 박스 + 점선 연결 (화살촉 없음)
# ══════════════════════════════════════════════════════════
re = cx + bw/2   # 박스 오른쪽 끝 = 8.0

annbox(axc, y2, aw, 0.72, [
    '1. F 라디칼 플럭스 모델  (CF4 포화 보정 포함)',
    '2. CFx 라디칼 플럭스 모델',
    '3. 이온 플럭스 모델  (CF4 fragment 이온화 보정)',
    '4. Sheath 이온 에너지 모델',
])
dseg(re + 0.05, y2, axc - aw/2 - 0.05, y2)

annbox(axc, y3, aw, 0.60, [
    '1. Clausing power-law  (수직 이온 전달)',
    '2. IAD acceptance-cone  (측면 이온 / shadow)',
    '3. Clausing × 지수감쇠  (중성 입자 전달)',
])
dseg(re + 0.05, y3, axc - aw/2 - 0.05, y3)

annbox(axc, y4, aw, 0.72, [
    '1. Bohdansky 스퍼터링 수율',
    '2. 이온강화 식각 인자  (threshold 모델)',
    '3. 수직 식각 속도  (화학 + 이온강화 + 스퍼터 - 패시)',
    '4. 측면 식각 속도  (화학 + IAD 이온강화 - 패시)',
])
dseg(re + 0.05, y4, axc - aw/2 - 0.05, y4)

annbox(axc, y5, aw, 0.60, [
    '1. 바닥/측벽 폴리머 동역학  (증착 - 이온제거)',
    '2. 마스크 개구부 진화  (Ar+스퍼터 + F화학 - CFx)',
    '3. Birth CD 모델  (IAD collimation 보정)',
])
dseg(re + 0.05, y5, axc - aw/2 - 0.05, y5)

# ══════════════════════════════════════════════════════════
prs.save(OUT)
print(f"Saved -> {OUT}")
print(f"슬라이드: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} 인치")
